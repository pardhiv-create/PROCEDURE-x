from pathlib import Path
import io
import pandas as pd

def parse_uploaded_file(filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    metadata = {"filename": filename, "extension": suffix, "bytes": len(content)}
    if suffix in {".txt", ".md", ".eml"}:
        text = content.decode("utf-8", errors="replace")
    elif suffix == ".csv":
        try:
            text = pd.read_csv(io.BytesIO(content)).to_string(index=False)
        except Exception:
            text = content.decode("utf-8", errors="replace")
    elif suffix in {".xlsx", ".xls"}:
        sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
        text = "\n\n".join(f"[Sheet: {name}]\n{df.to_string(index=False)}" for name, df in sheets.items())
        metadata["sheets"] = list(sheets.keys())
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = []
            for idx, page in enumerate(reader.pages, start=1):
                pages.append(f"[Page {idx}]\n{page.extract_text() or ''}")
            text = "\n\n".join(pages)
            metadata["pages"] = len(reader.pages)
            if not text.strip():
                raise ValueError("No selectable text found; PDF may be scanned. Run OCR first.")
        except ImportError as exc:
            raise RuntimeError("Install pypdf to parse PDF files.") from exc
    elif suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for index, table in enumerate(doc.tables, start=1):
            parts.append(f"[Table {index}]")
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    if not text.strip():
        raise ValueError("No text could be extracted from this document.")
    metadata["characters"] = len(text)
    return {"text": text[:30000], "metadata": metadata}
